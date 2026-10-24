"""Create small original Office comparison fixtures (openpyxl, python-pptx, Pillow)."""
from pathlib import Path
from io import BytesIO
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches
from PIL import Image

root = Path(__file__).resolve().parent
for side, changed in [('left', False), ('right', True)]:
    book = Workbook()
    sheet = book.active
    sheet.title = '売上'
    sheet.append(['商品', '数量', '単価', '合計'])
    sheet.append(['りんご', 4 if changed else 3, 120, '=B2*C2'])
    sheet.append(['みかん', 2, 80, '=B3*C3'])
    sheet['A1'].font = Font(bold=True, color='FFFFFF')
    sheet['A1'].fill = PatternFill('solid', fgColor='008080' if changed else '004080')
    sheet['C2'].number_format = '#,##0'
    sheet['A5'] = '複数行\nメモ & <test>'
    sheet.merge_cells('A6:C6')
    hidden = book.create_sheet('非表示の設定')
    hidden['B2'] = '新設定' if changed else '共通設定'
    hidden.sheet_state = 'hidden'
    if changed:
        book.create_sheet('追加シート')['A1'] = '追加'
    book.properties.creator = 'YY series'
    book.properties.created = datetime(2026, 1, 1)
    book.save(root / f'{side}.xlsx')

    deck = Presentation()
    for title in ['概要', '詳細']:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = '変更後の文章' if changed and title == '詳細' else '共通の文章\n日本語 & <test>'
        slide.notes_slide.notes_text_frame.text = '発表者ノート（変更）' if changed else '発表者ノート'
        image = BytesIO()
        Image.new('RGB', (12, 12), 'red' if changed else 'blue').save(image, format='PNG')
        image.seek(0)
        slide.shapes.add_picture(image, Inches(7), Inches(4), width=Inches(.5))
    if changed:
        ids = deck.slides._sldIdLst
        ids.insert(0, ids[1])
    deck.core_properties.author = 'YY series'
    deck.core_properties.created = datetime(2026, 1, 1)
    deck.save(root / f'{side}.pptx')
